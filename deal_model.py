"""
Independent Python implementation of the partnership deal model.

Purpose
  1. Cross-validate the Excel model (both read the same Inputs tab).
  2. Generate the scenario and tornado charts used in the README and memo.

Usage
  python deal_model.py [path/to/model.xlsx]
"""
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Partnership_Business_Case_Model.xlsx"
YEARS = 5
DRIVER_ROWS = {"subs": 5, "growth": 6, "hours": 7, "cpm": 8, "cannib": 9}
SCEN_COLS = {"Downside": "C", "Base": "D", "Upside": "E"}
GEN_ROWS = {"price": 12, "fee": 13, "admin": 14, "spots": 15, "fill": 16, "prod": 17, "esc": 18,
            "mkt": 19, "launch": 20, "rate": 21, "A_fee": 24, "A_esc": 25, "B_mg": 26,
            "B_share": 27, "C_share": 28, "upfront": 29, "w_down": 31, "w_base": 32, "w_up": 33,
            "sens_flex": 35}
LABELS = {"subs": "Year-1 subscribers", "growth": "Annual growth", "hours": "Year-1 viewing hours",
          "cpm": "Ad CPM", "cannib": "Cannibalization"}
NAMES = {"A": "A: Flat fee", "B": "B: MG + rev share", "C": "C: Rev share"}


def load_inputs(path):
    ws = load_workbook(path, data_only=True)["Inputs"]
    g = {k: ws[f"C{r}"].value for k, r in GEN_ROWS.items()}
    scen = {s: {k: ws[f"{c}{r}"].value for k, r in DRIVER_ROWS.items()} for s, c in SCEN_COLS.items()}
    return g, scen


def npv(rate, flows):
    """flows[0] = Year 0 (undiscounted); flows[1:] discounted at end of each year."""
    return flows[0] + sum(cf / (1 + rate) ** t for t, cf in enumerate(flows[1:], start=1))


def run(g, d):
    """Cash flows (Year 0..5) for each structure under one set of drivers."""
    cfs = {x: [-g["upfront"]] for x in "ABC"}
    subs, hours = d["subs"], d["hours"]
    for n in range(1, YEARS + 1):
        if n > 1:
            subs *= 1 + d["growth"]
            hours *= 1 + d["growth"]
        net_sub = subs * g["price"] * (1 - g["fee"])
        ad = hours * g["admin"] * g["spots"] * d["cpm"] / 1000 * g["fill"]
        deal_rev = net_sub + ad                      # revenue-share base
        incr = deal_rev - ad * d["cannib"]           # value to the platform
        costs = g["prod"] * (1 + g["esc"]) ** (n - 1) + incr * g["mkt"] + (g["launch"] if n == 1 else 0)
        contrib = incr - costs
        pay = {"A": g["A_fee"] * (1 + g["A_esc"]) ** (n - 1),
               "B": max(g["B_mg"], g["B_share"] * deal_rev),
               "C": g["C_share"] * deal_rev}
        for x in "ABC":
            cfs[x].append(contrib - pay[x])
    return cfs


def scenario_table(g, scen):
    out = {s: {x: npv(g["rate"], cf) for x, cf in run(g, d).items()} for s, d in scen.items()}
    w = {"Downside": g["w_down"], "Base": g["w_base"], "Upside": g["w_up"]}
    out["Prob-weighted"] = {x: sum(out[s][x] * w[s] for s in w) for x in "ABC"}
    return out


def tornado(g, scen, struct="B"):
    base = scen["Base"]
    base_npv = npv(g["rate"], run(g, base)[struct])
    rows = []
    for k in DRIVER_ROWS:
        res = []
        for sign in (-1, 1):
            d = dict(base)
            d[k] = base[k] * (1 + sign * g["sens_flex"])
            res.append(npv(g["rate"], run(g, d)[struct]) - base_npv)
        rows.append((k, res[0], res[1]))
    rows.sort(key=lambda r: abs(r[2] - r[1]))  # largest swing plotted on top
    return base_npv, rows


def validate_against_excel(path, table):
    """Largest absolute NPV difference between Python and the recalculated Excel model."""
    ws = load_workbook(path, data_only=True)["Scenarios"]
    cols = {"Downside": "B", "Base": "C", "Upside": "D", "Prob-weighted": "E"}
    return max(abs(ws[f"{c}{4 + i}"].value - table[s][x])
               for i, x in enumerate("ABC") for s, c in cols.items())


def charts(table, base_npv, rows, outdir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    scen = ["Downside", "Base", "Upside", "Prob-weighted"]
    fig, ax = plt.subplots(figsize=(8, 4.5))
    w = 0.25
    for i, x in enumerate("ABC"):
        ax.bar([j + (i - 1) * w for j in range(len(scen))], [table[s][x] for s in scen], w, label=NAMES[x])
    ax.axhline(0, color="black", lw=0.8)
    ax.set_xticks(range(len(scen)), scen)
    ax.set_ylabel("NPV ($mm)")
    ax.set_title("5-year NPV by deal structure and scenario")
    ax.legend(frameon=False)
    fig.tight_layout(); fig.savefig(outdir / "scenario_npv.png", dpi=160); plt.close(fig)

    fig, ax = plt.subplots(figsize=(8, 4))
    y = list(range(len(rows)))
    ax.barh(y, [r[1] for r in rows], color="#c0504d", label="Driver -20%")
    ax.barh(y, [r[2] for r in rows], color="#4f81bd", label="Driver +20%")
    ax.set_yticks(y, [LABELS[r[0]] for r in rows])
    ax.axvline(0, color="black", lw=0.8)
    ax.set_xlabel(f"Change in NPV vs. Base (${base_npv:,.0f}mm), Structure B")
    ax.set_title("Sensitivity: which assumptions move value most")
    ax.legend(frameon=False, loc="lower right")
    fig.tight_layout(); fig.savefig(outdir / "tornado.png", dpi=160); plt.close(fig)


if __name__ == "__main__":
    g, scen = load_inputs(XLSX)
    table = scenario_table(g, scen)
    print(f"{'NPV $mm':<14}" + "".join(f"{s:>15}" for s in table))
    for x in "ABC":
        print(f"Structure {x:<4}" + "".join(f"{table[s][x]:>15,.1f}" for s in table))
    base_npv, rows = tornado(g, scen)
    print("\nTornado (Structure B), largest swing last:")
    for k, lo, hi in rows:
        print(f"  {LABELS[k]:<22} -20%: {lo:>8,.1f}   +20%: {hi:>8,.1f}")
    try:
        diff = validate_against_excel(XLSX, table)
        print(f"\nValidation vs. Excel: max NPV difference = ${diff:.6f}mm -> {'PASS' if diff < 0.01 else 'FAIL'}")
    except Exception as e:
        print(f"\nValidation skipped (recalculate the Excel file first): {e}")
    out = ROOT
    charts(table, base_npv, rows, out)
    print(f"Charts written to {out}")
